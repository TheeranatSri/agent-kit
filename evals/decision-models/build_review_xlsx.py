"""Build local/review_decision_cases.xlsx for the user's label review of the decision-model test set.

Inputs: local/cases.jsonl (96 cases, company data, gitignored). No BQ / LLM / embedding calls.
Output: local/review_decision_cases.xlsx (gitignored). Model predictions are left out on purpose: the user's labels
must not be anchored on what a model answered.
Run: uv run --with openpyxl python build_review_xlsx.py
After the review: uv run --with openpyxl python build_review_xlsx.py --export  -> local/review_labels.csv
(case_id, user_label, user_note) for `run_eval.py --gold local/review_labels.csv`.
"""

from __future__ import annotations

import collections
import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
CASES = HERE / "local/cases.jsonl"
OUT = HERE / "local/review_decision_cases.xlsx"
LABELS_CSV = HERE / "local/review_labels.csv"
LANG_ORDER = {"th": 0, "en": 1, "mixed": 2}
TASK_TH = {
    "context_retention": "เก็บ output ของ tool ไว้ใน context ไหม",
    "agent_routing": "ส่งงานให้ agent ที่ถูกกว่าได้ไหม",
    "skill_selection": "ควรใช้ skill ไหน",
    "network_level": "job นี้ควรได้ network ระดับไหน",
}
HEAD = PatternFill("solid", fgColor="D9E1F2")
INPUT = PatternFill("solid", fgColor="FFF2CC")
GROUP = PatternFill("solid", fgColor="F2F2F2")


def scenario(case_id: str) -> str:
    m = re.match(r"(.+)_(th|en|mixed)_(\d+)$", case_id)
    return f"{m.group(1)}_{m.group(3)}" if m else case_id


def load_cases() -> list[dict]:
    rows = [json.loads(line) for line in CASES.open(encoding="utf-8") if line.strip()]
    rows.sort(key=lambda r: (r["task"], scenario(r["case_id"]), LANG_ORDER[r["language"]]))
    return rows


def build() -> None:
    cases = load_cases()
    assert len(cases) == 96, len(cases)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    counts = collections.Counter((c["task"], c["language"]) for c in cases)
    lines = [
        "ไฟล์รีวิว label ของชุดเทส decision model (OpenThai / Laya / กฎ deterministic)",
        f"สร้างเมื่อ {date.today()} จาก evals/decision-models/local/cases.jsonl ด้วย build_review_xlsx.py (ไม่มีการเรียก LLM / BQ)",
        "",
        "ทำไมต้องรีวิว: label ในไฟล์นี้ Codex เป็นคนเสนอ (proposed_label) เราไม่ใช้ LLM ตัดสิน ดังนั้น label ที่คุณยืนยันคือคำตอบที่ถูก",
        "ที่ใช้ให้คะแนน model ทุกตัว ถ้า label ผิด ผลเทียบ model ก็ผิดตาม",
        "",
        "วิธีใช้ sheet cases:",
        "1. อ่าน state (สถานการณ์), question และ options (label: ความหมาย)",
        "2. ช่อง user_label (สีเหลือง): เลือกจาก dropdown ถ้าเห็นด้วยกับ proposed_label ก็เลือกค่าเดียวกัน ถ้าไม่แน่ใจเลือก 'unsure'",
        "3. ช่อง user_note: เหตุผลสั้นๆ เมื่อไม่เห็นด้วย หรือเมื่อเคสเองมีปัญหา (เช่น ยังมีคำใบ้ / ไม่สมจริง)",
        "4. เคสเรียงเป็นกลุ่มละ 3 แถว (th, en, mixed) ของสถานการณ์เดียวกัน (คอลัมน์ scenario, แถบสีเทาสลับกลุ่ม)",
        "   ปกติ label ทั้ง 3 ภาษาควรเหมือนกัน ถ้าคุณให้ต่างกัน ช่วยเขียนเหตุผลใน user_note",
        "",
        "สิ่งที่ตั้งใจไม่ใส่: คำตอบของ Laya / baseline (เพื่อไม่ให้คำตอบของ model ชี้นำ label ของคุณ)",
        "ห้ามแก้ข้อความใน state / question / options; ถ้าเคสมีปัญหาให้เขียนใน user_note แทน",
        "",
        "ความหมายของ decision:",
    ] + [f"- {t}: {d}" for t, d in TASK_TH.items()] + [
        "",
        "สี: ฟ้า = หัวตาราง; เหลือง = ช่องที่คุณกรอก; เทา = แถบสลับกลุ่ม scenario (ไม่มีความหมายอื่น)",
        "",
        f"จำนวนเคส: {len(cases)} = 4 decision x 3 ภาษา x 8 (ตรวจแล้ว: ทุกช่อง = 8)",
        "เสร็จแล้ว: save ไฟล์ แล้วบอก Claude ให้ export -> local/review_labels.csv สำหรับให้คะแนน model",
    ]
    assert set(counts.values()) == {8}, counts
    for i, line in enumerate(lines, 1):
        readme.cell(i, 1, line)
    readme.column_dimensions["A"].width = 130

    ws = wb.create_sheet("cases")
    cols = ["case_id", "scenario", "task", "language", "difficulty", "state", "question", "options",
            "proposed_label", "rationale", "user_label", "user_note"]
    widths = [26, 22, 18, 9, 9, 70, 40, 45, 15, 50, 14, 35]
    for j, (name, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(1, j, name)
        cell.font = Font(bold=True)
        cell.fill = INPUT if name in ("user_label", "user_note") else HEAD
        ws.column_dimensions[cell.column_letter].width = w
    shade, last = False, None
    for i, c in enumerate(cases, 2):
        q = c["questions"][0]
        crit = q["request"]["criteria"]
        sc = scenario(c["case_id"])
        if sc != last:
            shade, last = not shade, sc
        values = [c["case_id"], sc, c["task"], c["language"], c.get("difficulty", ""), c["state"],
                  q["request"]["instructions"], "\n".join(f"{k}: {v}" for k, v in crit.items()),
                  q["expected"]["label"], c.get("rationale", ""), None, None]
        for j, v in enumerate(values, 1):
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if j >= 11:
                cell.fill = INPUT
            elif shade:
                cell.fill = GROUP
        options = ",".join(list(crit) + ["unsure"])
        assert len(options) < 255, (c["case_id"], options)  # Excel list validation limit
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"K{i}")
        assert q["expected"]["label"] in crit, c["case_id"]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{len(cases) + 1}"

    sm = wb.create_sheet("summary")
    sm.append(["task", "language", "cases", "filled", "agree_with_proposed"])
    for cell in sm[1]:
        cell.font, cell.fill = Font(bold=True), HEAD
    n = len(cases) + 1
    for t in TASK_TH:
        for lang in LANG_ORDER:
            r = sm.max_row + 1
            sm.append([t, lang, counts[(t, lang)],
                       f'=COUNTIFS(cases!$C$2:$C${n},A{r},cases!$D$2:$D${n},B{r},cases!$K$2:$K${n},"<>")',
                       f'=SUMPRODUCT((cases!$C$2:$C${n}=A{r})*(cases!$D$2:$D${n}=B{r})*(cases!$K$2:$K${n}=cases!$I$2:$I${n}))'])
    for col, w in zip("ABCDE", (20, 10, 8, 8, 22)):
        sm.column_dimensions[col].width = w
    wb.save(OUT)
    print(f"wrote {OUT} ({len(cases)} cases; per cell {sorted(set(counts.values()))})")


def export() -> None:
    ws = load_workbook(OUT, data_only=True)["cases"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    filled = [r for r in rows if r.get("user_label")]
    with LABELS_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["case_id", "user_label", "user_note"])
        for r in rows:
            # run_eval.py rejects labels outside the criteria: an 'unsure' case keeps the proposed label for scoring
            # and is listed in the note so the report can show it separately.
            unsure = r.get("user_label") == "unsure"
            label = "" if unsure else (r.get("user_label") or "")
            note = ("unsure; " if unsure else "") + (r.get("user_note") or "")
            w.writerow([r["case_id"], label, note])
    print(f"wrote {LABELS_CSV}: {len(filled)}/{len(rows)} labelled, "
          f"{sum(r['user_label'] == 'unsure' for r in filled)} unsure, "
          f"{sum(r['user_label'] not in (None, 'unsure', r['proposed_label']) for r in filled)} changed")


if __name__ == "__main__":
    export() if "--export" in sys.argv else build()
